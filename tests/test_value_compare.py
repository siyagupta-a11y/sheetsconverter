import sys
import unittest
from pathlib import Path

try:
    from openpyxl import Workbook, load_workbook
    HAS_OPENPYXL = True
except Exception:
    HAS_OPENPYXL = False

sys.path.insert(0, str(Path(__file__).resolve().parents[1] / "backend"))

from value_compare import (
    source_formula_values,
    compare_formula_values,
    apply_mismatch_highlights,
)


class ValueCompareTests(unittest.TestCase):
    def test_source_formula_values_extracts_effective(self):
        data = {
            "sheets": [
                {
                    "properties": {"title": "Output"},
                    "data": [
                        {
                            "startRow": 0,
                            "startColumn": 0,
                            "rowData": [
                                {
                                    "values": [
                                        {
                                            "userEnteredValue": {"formulaValue": "=1+1"},
                                            "effectiveValue": {"numberValue": 2},
                                        }
                                    ]
                                }
                            ],
                        }
                    ],
                }
            ]
        }
        m = source_formula_values(data)
        self.assertEqual(m["Output!A1"], 2)

    def test_compare_formula_values_detects_mismatch(self):
        src = {"Output!A1": 5, "Output!B1": 10}
        xls = {"Output!A1": 7, "Output!B1": 10}
        cmp = compare_formula_values(src, xls)
        self.assertEqual(cmp["checked_formula_cells"], 2)
        self.assertEqual(cmp["mismatch_count"], 1)
        self.assertEqual(cmp["mismatches"][0]["ref"], "Output!A1")

    @unittest.skipUnless(HAS_OPENPYXL, "openpyxl not available in local test environment")
    def test_apply_highlights_and_report_sheet(self):
        wb = Workbook()
        ws = wb.active
        ws.title = "Output"
        ws["A1"] = 1
        ws["B1"] = 2
        from io import BytesIO

        b = BytesIO()
        wb.save(b)
        out = apply_mismatch_highlights(
            b.getvalue(),
            [{"sheet": "Output", "cell": "A1", "source_value": 5, "excel_value": 7}],
        )
        wb2 = load_workbook(BytesIO(out))
        self.assertIn("⚠ Value Mismatches", wb2.sheetnames)
        self.assertEqual(wb2["⚠ Value Mismatches"]["A2"].value, "Output")


if __name__ == "__main__":
    unittest.main()
