import sys
import unittest
from pathlib import Path

try:
    from openpyxl import Workbook, load_workbook
    from openpyxl.worksheet.formula import ArrayFormula
    HAS_OPENPYXL = True
except Exception:
    HAS_OPENPYXL = False

sys.path.insert(0, str(Path(__file__).resolve().parents[1] / "backend"))

from excel_formula_patcher import patch_uploaded_workbook
from excel_formula_patcher import _extract_formula, _restore_formula


class ExcelFormulaHelpersTests(unittest.TestCase):
    def test_extract_formula_normal_and_array(self):
        f, kind = _extract_formula("=SUM(A1:A3)")
        self.assertEqual(f, "=SUM(A1:A3)")
        self.assertEqual(kind, "normal")

        f2, kind2 = _extract_formula("{=SUM(A1:A3)}")
        self.assertEqual(f2, "=SUM(A1:A3)")
        self.assertEqual(kind2, "array")

    def test_restore_formula(self):
        self.assertEqual(_restore_formula("=A1", "normal"), "=A1")
        self.assertEqual(_restore_formula("=A1", "array"), "=A1")

    @unittest.skipUnless(HAS_OPENPYXL, "openpyxl not available in local test environment")
    def test_extract_and_restore_array_formula_object(self):
        af = ArrayFormula("C1:C2", "=IF(1,2,3)")
        formula, kind = _extract_formula(af)
        self.assertEqual(formula, "=IF(1,2,3)")
        self.assertEqual(kind, ("array_obj", "C1:C2"))

        restored = _restore_formula("=IF(1,9,3)", kind)
        self.assertEqual(restored.ref, "C1:C2")
        self.assertEqual(restored.text, "=IF(1,9,3)")


@unittest.skipUnless(HAS_OPENPYXL, "openpyxl not available in local test environment")
class ExcelFormulaPatcherTests(unittest.TestCase):
    def test_patch_uploaded_workbook_rewrites_formula_and_adds_notes(self):
        from io import BytesIO

        wb = Workbook()
        ws = wb.active
        ws.title = "Model"
        ws["A1"] = '=IFERROR(1/0)'
        ws["B1"] = '=QUERY(A1:A5,"select A")'

        buf = BytesIO()
        wb.save(buf)

        out_bytes, warnings = patch_uploaded_workbook(buf.getvalue())
        wb2 = load_workbook(BytesIO(out_bytes))

        self.assertEqual(wb2["Model"]["A1"].value, '=IFERROR(1/0, "")')
        self.assertEqual(wb2["Model"]["B1"].value, '=QUERY(A1:A5,"select A")')
        self.assertTrue(any(w.get("has_unsupported") for w in warnings))
        self.assertIn("⚠ Conversion Notes", wb2.sheetnames)


if __name__ == "__main__":
    unittest.main()
