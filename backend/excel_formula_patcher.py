import io

from formula_converter import convert_formula

try:
    from openpyxl.worksheet.formula import ArrayFormula
except Exception:  # pragma: no cover - openpyxl may be unavailable in some test envs
    ArrayFormula = None


def _extract_formula(value: str):
    """
    Return (formula, kind) where formula starts with '=' and kind is:
    - 'normal' for =...
    - 'array' for {=...} string form
    - ('array_obj', ref) for openpyxl ArrayFormula objects
    Returns (None, None) if value is not a recognized formula string.
    """
    if not isinstance(value, str):
        if ArrayFormula is not None and isinstance(value, ArrayFormula):
            text = (value.text or "").strip()
            if text.startswith("="):
                return text, ("array_obj", value.ref)
        return None, None
    s = value.strip()
    if s.startswith("="):
        return s, "normal"
    if s.startswith("{=") and s.endswith("}"):
        return s[1:-1], "array"
    return None, None


def _restore_formula(formula: str, kind: str):
    if isinstance(kind, tuple) and len(kind) == 2 and kind[0] == "array_obj":
        ref = kind[1]
        if ArrayFormula is not None:
            return ArrayFormula(ref, formula)
        return formula

    # Never write brace-wrapped strings back into cell.value. Excel stores
    # formulas as "=..." and handles array semantics separately.
    # Writing "{=...}" as text can trigger #NAME? in output workbooks.
    return formula


def _write_notes_sheet(wb, warnings: list):
    from openpyxl.styles import Font
    from openpyxl.utils import get_column_letter

    notes_name = "⚠ Conversion Notes"
    if notes_name in wb.sheetnames:
        del wb[notes_name]
    if not warnings:
        return

    ws_notes = wb.create_sheet(title=notes_name)
    col_widths = [12, 8, 45, 45, 55]
    headers = ["Sheet", "Cell", "Original Formula", "Converted Formula", "Notes"]
    for i, (header, width) in enumerate(zip(headers, col_widths), 1):
        c = ws_notes.cell(row=1, column=i, value=header)
        c.font = Font(bold=True)
        ws_notes.column_dimensions[get_column_letter(i)].width = width

    for row_idx, w in enumerate(warnings, 2):
        ws_notes.cell(row=row_idx, column=1, value=w.get("sheet", ""))
        ws_notes.cell(row=row_idx, column=2, value=w.get("cell", ""))
        ws_notes.cell(row=row_idx, column=3, value=w.get("original", ""))
        converted = w.get("converted", "")
        ws_notes.cell(row=row_idx, column=4, value=converted if converted else "(cleared)")
        ws_notes.cell(row=row_idx, column=5, value="; ".join(w.get("warnings", [])))


def patch_uploaded_workbook(xlsx_bytes: bytes) -> tuple[bytes, list]:
    from openpyxl import load_workbook

    wb = load_workbook(io.BytesIO(xlsx_bytes))
    warnings = []

    for ws in wb.worksheets:
        if ws.title == "⚠ Conversion Notes":
            continue
        for row in ws.iter_rows():
            for cell in row:
                formula, kind = _extract_formula(cell.value)
                if not formula:
                    continue

                original = cell.value
                res = convert_formula(formula)

                if res.formula == "":
                    cell.value = None
                else:
                    cell.value = _restore_formula(res.formula, kind)

                if res.warnings:
                    warnings.append(
                        {
                            "sheet": ws.title,
                            "cell": cell.coordinate,
                            "original": original,
                            "converted": res.formula,
                            "warnings": res.warnings,
                            "has_unsupported": res.has_unsupported,
                        }
                    )

    _write_notes_sheet(wb, warnings)
    out = io.BytesIO()
    wb.save(out)
    return out.getvalue(), warnings
