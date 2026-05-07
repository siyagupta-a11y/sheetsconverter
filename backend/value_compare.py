import io
import os
import shutil
import subprocess
import tempfile
from typing import Any

MISMATCH_FILL_RGB = "FFFFC7CE"


def _col_to_letter(col: int) -> str:
    result = ""
    while col > 0:
        col, rem = divmod(col - 1, 26)
        result = chr(65 + rem) + result
    return result


def _effective_value(cell: dict):
    ev = cell.get("effectiveValue", {})
    if "numberValue" in ev:
        return ev["numberValue"]
    if "stringValue" in ev:
        return ev["stringValue"]
    if "boolValue" in ev:
        return ev["boolValue"]
    if "errorValue" in ev:
        err = ev.get("errorValue", {})
        t = err.get("type") or "ERROR"
        return f"#{t}"
    return None


def source_formula_values(spreadsheet_data: dict) -> dict:
    """Map formula cells to Google effective values: {'Sheet!A1': value}."""
    out = {}
    for sheet in spreadsheet_data.get("sheets", []):
        sheet_title = sheet.get("properties", {}).get("title", "Sheet")
        for data_range in sheet.get("data", []):
            start_row = data_range.get("startRow", 0)
            start_col = data_range.get("startColumn", 0)
            for r_idx, row in enumerate(data_range.get("rowData", [])):
                for c_idx, cell in enumerate(row.get("values", [])):
                    formula = cell.get("userEnteredValue", {}).get("formulaValue")
                    if not formula:
                        continue
                    row_1 = start_row + r_idx + 1
                    col_1 = start_col + c_idx + 1
                    key = f"{sheet_title}!{_col_to_letter(col_1)}{row_1}"
                    out[key] = _effective_value(cell)
    return out


def workbook_values_from_bytes(xlsx_bytes: bytes, data_only: bool = True) -> dict:
    from openpyxl import load_workbook

    wb = load_workbook(io.BytesIO(xlsx_bytes), data_only=data_only)
    out = {}
    for ws in wb.worksheets:
        for row in ws.iter_rows():
            for cell in row:
                if cell.value is None:
                    continue
                out[f"{ws.title}!{cell.coordinate}"] = cell.value
    return out


def _find_office_binary():
    return shutil.which("libreoffice") or shutil.which("soffice")


def evaluate_excel_values(xlsx_bytes: bytes) -> dict:
    """
    Best-effort Excel-style evaluation via LibreOffice headless.
    Returns {'ok': bool, 'engine': str, 'values': dict, 'note': str}.
    """
    binary = _find_office_binary()
    if not binary:
        return {"ok": False, "engine": "unavailable", "values": {}, "note": "LibreOffice not installed"}

    with tempfile.TemporaryDirectory() as td:
        in_path = os.path.join(td, "input.xlsx")
        out_path = os.path.join(td, "input.xlsx")
        with open(in_path, "wb") as f:
            f.write(xlsx_bytes)

        try:
            subprocess.run(
                [binary, "--headless", "--convert-to", "xlsx", "--outdir", td, in_path],
                check=True,
                stdout=subprocess.PIPE,
                stderr=subprocess.PIPE,
                timeout=120,
            )
            if not os.path.exists(out_path):
                # Some builds can emit uppercase extension.
                alt = os.path.join(td, "input.XLSX")
                if os.path.exists(alt):
                    out_path = alt
            with open(out_path, "rb") as f:
                recalced = f.read()
            return {
                "ok": True,
                "engine": "libreoffice",
                "values": workbook_values_from_bytes(recalced, data_only=True),
                "note": "",
            }
        except Exception as e:
            return {
                "ok": False,
                "engine": "unavailable",
                "values": {},
                "note": f"Evaluation failed: {str(e)[:180]}",
            }


def _is_number(v: Any) -> bool:
    return isinstance(v, (int, float)) and not isinstance(v, bool)


def _equivalent(a: Any, b: Any, tol: float = 1e-9) -> bool:
    if a is None and b is None:
        return True
    if _is_number(a) and _is_number(b):
        aa = float(a)
        bb = float(b)
        diff = abs(aa - bb)
        lim = max(tol, tol * max(abs(aa), abs(bb), 1.0))
        return diff <= lim
    return a == b


def compare_formula_values(source_values: dict, excel_values: dict) -> dict:
    mismatches = []
    checked = 0
    for ref, src in source_values.items():
        checked += 1
        xls = excel_values.get(ref)
        if not _equivalent(src, xls):
            sheet, cell = ref.split("!", 1)
            mismatches.append(
                {
                    "ref": ref,
                    "sheet": sheet,
                    "cell": cell,
                    "source_value": src,
                    "excel_value": xls,
                }
            )
    return {
        "checked_formula_cells": checked,
        "mismatch_count": len(mismatches),
        "mismatches": mismatches,
    }


def apply_mismatch_highlights(xlsx_bytes: bytes, mismatches: list) -> bytes:
    from openpyxl import load_workbook
    from openpyxl.styles import PatternFill
    from openpyxl.utils import get_column_letter

    if not mismatches:
        return xlsx_bytes

    mismatch_fill = PatternFill(fill_type="solid", fgColor=MISMATCH_FILL_RGB)
    wb = load_workbook(io.BytesIO(xlsx_bytes))
    for m in mismatches:
        sheet = m.get("sheet")
        cell = m.get("cell")
        if not sheet or not cell or sheet not in wb.sheetnames:
            continue
        ws = wb[sheet]
        ws[cell].fill = mismatch_fill

    sheet_name = "⚠ Value Mismatches"
    if sheet_name in wb.sheetnames:
        del wb[sheet_name]
    ws = wb.create_sheet(title=sheet_name)
    headers = ["Sheet", "Cell", "Sheets value", "Excel value"]
    widths = [18, 10, 25, 25]
    for i, (h, w) in enumerate(zip(headers, widths), 1):
        ws.cell(row=1, column=i, value=h)
        ws.column_dimensions[get_column_letter(i)].width = w

    for idx, m in enumerate(mismatches, 2):
        ws.cell(row=idx, column=1, value=m.get("sheet"))
        ws.cell(row=idx, column=2, value=m.get("cell"))
        ws.cell(row=idx, column=3, value=str(m.get("source_value")))
        ws.cell(row=idx, column=4, value=str(m.get("excel_value")))

    out = io.BytesIO()
    wb.save(out)
    return out.getvalue()
