import io
from typing import Optional
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.styles.colors import Color
from openpyxl.utils import get_column_letter
from openpyxl.comments import Comment
from openpyxl.workbook.defined_name import DefinedName
from openpyxl.utils.cell import quote_sheetname


def _to_argb(color_obj: dict, default: str = "FF000000") -> str:
    """Google Sheets RGBA floats → 8-char ARGB hex for openpyxl."""
    if not color_obj:
        return default
    r = int(color_obj.get('red', 0) * 255)
    g = int(color_obj.get('green', 0) * 255)
    b = int(color_obj.get('blue', 0) * 255)
    a = int(color_obj.get('alpha', 1.0) * 255)
    return f"{a:02X}{r:02X}{g:02X}{b:02X}"


def _resolve_color(color_field) -> dict:
    """Accept either a direct rgbColor dict or a colorStyle wrapper."""
    if not color_field:
        return {}
    if 'rgbColor' in color_field:
        return color_field['rgbColor']
    # Direct dict with 'red'/'green'/'blue' keys
    if 'red' in color_field or 'green' in color_field or 'blue' in color_field:
        return color_field
    return {}


BORDER_STYLES = {
    'DOTTED': 'dotted',
    'DASHED': 'dashed',
    'SOLID': 'thin',
    'SOLID_MEDIUM': 'medium',
    'SOLID_THICK': 'thick',
    'DOUBLE': 'double',
}

HALIGN = {'LEFT': 'left', 'CENTER': 'center', 'RIGHT': 'right', 'JUSTIFY': 'justify'}
VALIGN = {'TOP': 'top', 'MIDDLE': 'center', 'BOTTOM': 'bottom'}

NUMBER_FORMAT_DEFAULTS = {
    'TEXT': '@',
    'NUMBER': '0.##########',
    'PERCENT': '0.00%',
    'CURRENCY': '"$"#,##0.00',
    'DATE': 'MM/DD/YYYY',
    'TIME': 'HH:MM:SS AM/PM',
    'DATE_TIME': 'MM/DD/YYYY HH:MM:SS',
    'SCIENTIFIC': '0.00E+00',
}


def _make_font(text_format: dict) -> Optional[Font]:
    if not text_format:
        return None
    fg = _resolve_color(
        text_format.get('foregroundColorStyle') or text_format.get('foregroundColor')
    )
    argb = _to_argb(fg, "FF000000") if fg else "FF000000"
    return Font(
        bold=text_format.get('bold', False),
        italic=text_format.get('italic', False),
        underline='single' if text_format.get('underline') else None,
        strike=text_format.get('strikethrough', False),
        size=text_format.get('fontSize') or None,
        name=text_format.get('fontFamily') or None,
        color=Color(rgb=argb),
    )


def _make_fill(fmt: dict) -> Optional[PatternFill]:
    bg = _resolve_color(fmt.get('backgroundColorStyle') or fmt.get('backgroundColor'))
    if not bg:
        return None
    r, g, b = bg.get('red', 1.0), bg.get('green', 1.0), bg.get('blue', 1.0)
    if r == 1.0 and g == 1.0 and b == 1.0:
        return None  # default white background, skip
    argb = _to_argb(bg)
    return PatternFill(fill_type='solid', fgColor=Color(rgb=argb))


def _make_border(borders: dict) -> Optional[Border]:
    if not borders:
        return None

    def side(data):
        if not data:
            return Side(style=None)
        style = BORDER_STYLES.get(data.get('style', 'NONE'))
        if not style:
            return Side(style=None)
        color = _resolve_color(data.get('colorStyle') or data.get('color'))
        argb = _to_argb(color, "FF000000") if color else "FF000000"
        return Side(style=style, color=Color(rgb=argb))

    return Border(
        left=side(borders.get('left')),
        right=side(borders.get('right')),
        top=side(borders.get('top')),
        bottom=side(borders.get('bottom')),
    )


def _make_alignment(fmt: dict) -> Optional[Alignment]:
    if not fmt:
        return None
    h = HALIGN.get(fmt.get('horizontalAlignment', ''))
    v = VALIGN.get(fmt.get('verticalAlignment', ''))
    wrap = fmt.get('wrapStrategy') == 'WRAP'
    if not h and not v and not wrap:
        return None
    return Alignment(horizontal=h, vertical=v, wrap_text=wrap)


def _number_format(fmt: dict) -> Optional[str]:
    nf = fmt.get('numberFormat', {})
    if not nf:
        return None
    if nf.get('pattern'):
        return nf['pattern']
    return NUMBER_FORMAT_DEFAULTS.get(nf.get('type', ''))


def _col_to_letter(col: int) -> str:
    result = ''
    while col > 0:
        col, rem = divmod(col - 1, 26)
        result = chr(65 + rem) + result
    return result


def _grid_range_to_a1(grid_range: dict, sheet_title: str):
    if not grid_range:
        return None
    if 'startRowIndex' not in grid_range or 'endRowIndex' not in grid_range:
        return None
    if 'startColumnIndex' not in grid_range or 'endColumnIndex' not in grid_range:
        return None

    r1 = int(grid_range['startRowIndex']) + 1
    r2 = int(grid_range['endRowIndex'])
    c1 = int(grid_range['startColumnIndex']) + 1
    c2 = int(grid_range['endColumnIndex'])
    if r2 < r1 or c2 < c1:
        return None
    return f"{quote_sheetname(sheet_title)}!{_col_to_letter(c1)}{r1}:{_col_to_letter(c2)}{r2}"


def _apply_sheet_properties(ws, sheet: dict):
    props = sheet.get('properties', {})
    grid = props.get('gridProperties', {})
    frozen_rows = int(grid.get('frozenRowCount', 0) or 0)
    frozen_cols = int(grid.get('frozenColumnCount', 0) or 0)
    if frozen_rows > 0 or frozen_cols > 0:
        ws.freeze_panes = ws.cell(row=frozen_rows + 1, column=frozen_cols + 1)

    if props.get('hidden'):
        ws.sheet_state = 'hidden'

    tab = props.get('tabColorStyle') or props.get('tabColor')
    if tab:
        rgb = _resolve_color(tab)
        if rgb:
            ws.sheet_properties.tabColor = _to_argb(rgb, default="FF4B5563")

    basic_filter = sheet.get('basicFilter', {})
    flt_range = basic_filter.get('range')
    if flt_range:
        r1 = int(flt_range.get('startRowIndex', 0)) + 1
        r2 = int(flt_range.get('endRowIndex', 0))
        c1 = int(flt_range.get('startColumnIndex', 0)) + 1
        c2 = int(flt_range.get('endColumnIndex', 0))
        if r2 >= r1 and c2 >= c1:
            ws.auto_filter.ref = f"{_col_to_letter(c1)}{r1}:{_col_to_letter(c2)}{r2}"


def _apply_named_ranges(wb: Workbook, spreadsheet_data: dict):
    sheet_by_id = {}
    for sh in spreadsheet_data.get('sheets', []):
        props = sh.get('properties', {})
        sid = props.get('sheetId')
        title = props.get('title', 'Sheet')
        sheet_by_id[sid] = title[:31]

    for nr in spreadsheet_data.get('namedRanges', []):
        name = nr.get('name')
        rng = nr.get('range')
        if not name or not rng:
            continue
        sheet_title = sheet_by_id.get(rng.get('sheetId'))
        if not sheet_title:
            continue
        a1 = _grid_range_to_a1(rng, sheet_title)
        if not a1:
            continue
        try:
            wb.defined_names.add(DefinedName(name=name, attr_text=a1))
        except Exception:
            pass


def build_excel(spreadsheet_data: dict, warnings: list) -> bytes:
    wb = Workbook()
    wb.remove(wb.active)

    for sheet in spreadsheet_data.get('sheets', []):
        props = sheet.get('properties', {})
        title = props.get('title', 'Sheet')[:31]  # Excel limit

        ws = wb.create_sheet(title=title)
        _apply_sheet_properties(ws, sheet)

        for data_range in sheet.get('data', []):
            start_row = data_range.get('startRow', 0)
            start_col = data_range.get('startColumn', 0)

            for i, col_meta in enumerate(data_range.get('columnMetadata', [])):
                px = col_meta.get('pixelSize')
                if px:
                    ws.column_dimensions[get_column_letter(start_col + i + 1)].width = px / 7.0

            for i, row_meta in enumerate(data_range.get('rowMetadata', [])):
                px = row_meta.get('pixelSize')
                if px:
                    ws.row_dimensions[start_row + i + 1].height = px * 0.75

            for row_idx, row_data in enumerate(data_range.get('rowData', [])):
                abs_row = start_row + row_idx + 1
                for col_idx, cell_data in enumerate(row_data.get('values', [])):
                    abs_col = start_col + col_idx + 1
                    cell = ws.cell(row=abs_row, column=abs_col)

                    uev = cell_data.get('userEnteredValue', {})
                    if 'formulaValue' in uev and uev['formulaValue']:
                        cell.value = uev['formulaValue']
                    elif 'numberValue' in uev:
                        cell.value = uev['numberValue']
                    elif 'stringValue' in uev:
                        cell.value = uev['stringValue']
                    elif 'boolValue' in uev:
                        cell.value = uev['boolValue']

                    if cell_data.get('hyperlink'):
                        cell.hyperlink = cell_data.get('hyperlink')
                    if cell_data.get('note'):
                        cell.comment = Comment(cell_data['note'], "Sheets")

                    fmt = cell_data.get('userEnteredFormat', {})
                    if not fmt:
                        continue

                    font = _make_font(fmt.get('textFormat', {}))
                    if font:
                        cell.font = font

                    fill = _make_fill(fmt)
                    if fill:
                        cell.fill = fill

                    border = _make_border(fmt.get('borders', {}))
                    if border:
                        cell.border = border

                    alignment = _make_alignment(fmt)
                    if alignment:
                        cell.alignment = alignment

                    nf = _number_format(fmt)
                    if nf:
                        cell.number_format = nf

        for merge in sheet.get('merges', []):
            r1 = merge.get('startRowIndex', 0) + 1
            r2 = merge.get('endRowIndex', 0)
            c1 = merge.get('startColumnIndex', 0) + 1
            c2 = merge.get('endColumnIndex', 0)
            if r2 >= r1 and c2 >= c1:
                try:
                    ws.merge_cells(start_row=r1, start_column=c1, end_row=r2, end_column=c2)
                except Exception:
                    pass

    _apply_named_ranges(wb, spreadsheet_data)

    if warnings:
        ws_notes = wb.create_sheet(title='⚠ Conversion Notes')
        col_widths = [12, 8, 45, 45, 55]
        headers = ['Sheet', 'Cell', 'Original Formula', 'Converted Formula', 'Notes']
        for i, (header, width) in enumerate(zip(headers, col_widths), 1):
            c = ws_notes.cell(row=1, column=i, value=header)
            c.font = Font(bold=True)
            ws_notes.column_dimensions[get_column_letter(i)].width = width

        for row_idx, w in enumerate(warnings, 2):
            ws_notes.cell(row=row_idx, column=1, value=w.get('sheet', ''))
            ws_notes.cell(row=row_idx, column=2, value=w.get('cell', ''))
            ws_notes.cell(row=row_idx, column=3, value=w.get('original', ''))
            converted = w.get('converted', '')
            ws_notes.cell(row=row_idx, column=4, value=converted if converted else '(cleared)')
            ws_notes.cell(row=row_idx, column=5, value='; '.join(w.get('warnings', [])))

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def patch_workbook_formulas(xlsx_bytes: bytes, spreadsheet_data: dict, warnings: list) -> bytes:
    """
    Patch formulas in a native-exported workbook with converted formulas from spreadsheet_data.
    This preserves Google's export fidelity while applying compatibility rewrites.
    """
    wb = load_workbook(io.BytesIO(xlsx_bytes))

    for sheet in spreadsheet_data.get('sheets', []):
        title = sheet.get('properties', {}).get('title', 'Sheet')[:31]
        if title not in wb.sheetnames:
            continue
        ws = wb[title]

        for data_range in sheet.get('data', []):
            start_row = data_range.get('startRow', 0)
            start_col = data_range.get('startColumn', 0)
            for row_idx, row in enumerate(data_range.get('rowData', [])):
                for col_idx, cell_data in enumerate(row.get('values', [])):
                    uev = cell_data.get('userEnteredValue', {})
                    if 'formulaValue' not in uev:
                        continue
                    row_1 = start_row + row_idx + 1
                    col_1 = start_col + col_idx + 1
                    ws.cell(row=row_1, column=col_1).value = uev.get('formulaValue')

    if warnings:
        notes_name = '⚠ Conversion Notes'
        if notes_name in wb.sheetnames:
            del wb[notes_name]
        ws_notes = wb.create_sheet(title=notes_name)
        col_widths = [12, 8, 45, 45, 55]
        headers = ['Sheet', 'Cell', 'Original Formula', 'Converted Formula', 'Notes']
        for i, (header, width) in enumerate(zip(headers, col_widths), 1):
            c = ws_notes.cell(row=1, column=i, value=header)
            c.font = Font(bold=True)
            ws_notes.column_dimensions[get_column_letter(i)].width = width
        for row_idx, w in enumerate(warnings, 2):
            ws_notes.cell(row=row_idx, column=1, value=w.get('sheet', ''))
            ws_notes.cell(row=row_idx, column=2, value=w.get('cell', ''))
            ws_notes.cell(row=row_idx, column=3, value=w.get('original', ''))
            converted = w.get('converted', '')
            ws_notes.cell(row=row_idx, column=4, value=converted if converted else '(cleared)')
            ws_notes.cell(row=row_idx, column=5, value='; '.join(w.get('warnings', [])))

    out = io.BytesIO()
    wb.save(out)
    return out.getvalue()
